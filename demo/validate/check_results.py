#!/usr/bin/env python3
"""
check_results.py — Validation assertions for the Verus models demo.

Three check groups (run after the pipeline):
  1. DeprRateAnalyzR: the R models must RECOVER the known ground-truth
     parameters the synthetic data was generated with (proof the statistical
     pipeline works, not just that it runs).
  2. NACody extraction: every field text-mined from the sample PDFs must
     exactly match the ground truth the PDFs were generated from.
  3. Dashboard: workbook structure + Model_Params consistency with R output.

Exit code 0 = all pass.
"""
import csv
import json
import sys
from pathlib import Path

HERE = Path(__file__).parent
DATA = HERE.parent / "data"
R_OUT = HERE.parent / "deprrate" / "output"
NACODY = HERE.parent / "nacody"

FAILURES = []


def check(label, ok, detail=""):
    status = "PASS" if ok else "FAIL"
    print(f"  [{status}] {label}" + (f"  ({detail})" if detail else ""))
    if not ok:
        FAILURES.append(label)


def read_csv(path):
    with open(path) as f:
        return list(csv.DictReader(f))


def close(est, true, tol):
    return abs(est - true) <= tol


def main():
    print("== 1. DeprRateAnalyzR: ground-truth parameter recovery ==")
    truth = json.load(open(DATA / "ground_truth_params.json"))["categories"]
    for cat, p in truth.items():
        coefs = {r["Term"]: float(r["Estimate"])
                 for r in read_csv(R_OUT / f"coefficients_{cat.replace(' ', '_')}.csv")}
        check(f"{cat}: Age coefficient recovered",
              close(coefs["Age"], p["b_age"], 0.02),
              f"est={coefs['Age']:.4f} true={p['b_age']}")
        check(f"{cat}: LogNRC coefficient recovered",
              close(coefs["LogNRC"], p["b_lognrc"], 0.10),
              f"est={coefs['LogNRC']:.4f} true={p['b_lognrc']}")
        usage = next((v for k, v in coefs.items() if "Km/" in k or "Hours/" in k), None)
        if p["b_usage"] != 0:
            check(f"{cat}: usage coefficient recovered",
                  usage is not None and close(usage, p["b_usage"], 0.01),
                  f"est={usage and round(usage,4)} true={p['b_usage']}")
        if p["b_rebuilt"] != 0:
            check(f"{cat}: rebuilt-title discount recovered",
                  close(coefs["Title_StatusRebuilt"], p["b_rebuilt"], 0.06),
                  f"est={coefs['Title_StatusRebuilt']:.4f} true={p['b_rebuilt']}")
        if p.get("logit"):
            lg = {r["Term"]: float(r["Estimate"])
                  for r in read_csv(R_OUT / "logit_title_status.csv")}
            check(f"{cat}: logit Age coefficient recovered",
                  close(lg["Age"], p["logit"]["g_age"], 0.10),
                  f"est={lg['Age']:.4f} true={p['logit']['g_age']}")

    depr = {r["Sub_Category"]: r for r in read_csv(R_OUT / "depr_rates_demo.csv")}
    for cat, p in truth.items():
        import math
        true_rate = 1 - math.exp(p["b_age"])
        check(f"{cat}: annual depreciation rate plausible",
              close(float(depr[cat]["AnnualDeprRate"]), true_rate, 0.02),
              f"est={depr[cat]['AnnualDeprRate']} true={true_rate:.4f}")

    print("== 2. NACody: PDF extraction exact-match ==")
    gt = json.load(open(NACODY / "sample_reports" / "ground_truth.json"))
    extracted = {r["Workfile"]: r for r in read_csv(NACODY / "extracted_assets.csv")}
    check("all sample reports extracted", set(gt) == set(extracted),
          f"{len(extracted)}/{len(gt)}")
    for wf, want in gt.items():
        got = extracted.get(wf, {})
        for field in ["PublishedDate", "PublishedYear", "Customer", "Client",
                      "Item", "SerialNumber", "FMV_Low", "FMV_High",
                      "OLV_Low", "OLV_High", "Condition"]:
            check(f"{wf}.{field}", got.get(field, "") == str(want[field]),
                  f"got={got.get(field)!r} want={want[field]!r}")
        check(f"{wf}: all spec bullets captured",
              int(got.get("NumSpecs", 0)) == want["NumSpecs"])
        check(f"{wf}: all market comps captured",
              int(got.get("NumComps", 0)) == want["NumComps"])

    print("== 3. Dashboard: structure + coefficient consistency ==")
    from openpyxl import load_workbook
    xlsx = NACODY / "NACody_Demo_Dashboard.xlsx"
    check("dashboard file exists", xlsx.exists())
    if xlsx.exists():
        wb = load_workbook(xlsx)
        for sheet in ["README", "Estimator", "Model_Params", "Depr_Rates",
                      "Market_Comps", "Client_Summary"]:
            check(f"sheet present: {sheet}", sheet in wb.sheetnames)
        mp = wb["Model_Params"]
        params = {row[0]: row for row in mp.iter_rows(min_row=2, values_only=True)}
        for cat in truth:
            coefs = {r["Term"]: float(r["Estimate"])
                     for r in read_csv(R_OUT / f"coefficients_{cat.replace(' ', '_')}.csv")}
            check(f"Model_Params[{cat}] matches R output",
                  cat in params and close(params[cat][2], coefs["Age"], 1e-4))
        est = wb["Estimator"]
        check("Estimator has live VLOOKUP formula",
              isinstance(est["B10"].value, str) and "VLOOKUP" in est["B10"].value)

    print()
    if FAILURES:
        print(f"VALIDATION FAILED: {len(FAILURES)} check(s):")
        for f in FAILURES:
            print(f"  - {f}")
        sys.exit(1)
    print("ALL CHECKS PASSED.")


if __name__ == "__main__":
    main()
