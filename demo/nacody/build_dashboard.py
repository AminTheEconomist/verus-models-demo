#!/usr/bin/env python3
"""
build_dashboard.py — Assemble the NACody-style Excel dashboard from the demo
pipeline outputs.

Reproduces the real R -> Excel workflow evidenced by the surviving regression
workbook and market-research worksheet, on synthetic data only:

  Sheets:
    README           what this workbook is, in one screen
    Estimator        enter Category/Age/Usage/NRC -> live FMV estimate with
                     Low/High band (formulas reference Model_Params)
    Model_Params     per-category coefficients exported by depr_rate_analyzer.R
    Depr_Rates       the depreciation-rates table (R output)
    Market_Comps     comparables text-mined from the sample PDFs
    Client_Summary   per-client rollup in the worksheet layout (fake client)

Usage:  python3 build_dashboard.py [r_outdir] [extracted_csv] [out_xlsx]
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
R_OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / "deprrate" / "output"
EXTRACTED = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / "extracted_assets.csv"
OUT_XLSX = Path(sys.argv[3]) if len(sys.argv) > 3 else HERE / "NACody_Demo_Dashboard.xlsx"

CATEGORIES = ["Highway Tractor", "Crawler Excavator", "Dry Van Trailer"]

HEAD = Font(bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor="1F4E5F")
TITLE = Font(bold=True, size=14)


def read_csv(path):
    with open(path) as f:
        return list(csv.DictReader(f))


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEAD
        cell.fill = HEADFILL


def main():
    wb = Workbook()

    # --- README ---------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    ws["A1"] = "NACody Demo Dashboard (synthetic data)"
    ws["A1"].font = TITLE
    for i, line in enumerate([
        "All data in this workbook is synthetic - no Verus client data.",
        "",
        "Pipeline: synthetic asset master -> R (depr_rate_analyzer.R) fits",
        "log-linear depreciation models per asset class -> coefficients land",
        "in Model_Params -> the Estimator sheet prices an asset live with",
        "formulas, exactly the RegResults.xlsx pattern used at Verus.",
        "",
        "Market_Comps is text-mined from sample appraisal PDFs with",
        "extract_market_research.py (the ByCody/PdfScrape pattern).",
        "",
        "Try it: open Estimator, change Age or Usage, watch the FMV band move.",
    ], start=3):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 80

    # --- Model_Params (from R coefficient CSVs) --------------------------
    ws = wb.create_sheet("Model_Params")
    cols = ["Category", "Intercept", "B_Age", "B_Usage", "B_LogNRC", "B_Rebuilt", "ResidualSE"]
    ws.append(cols)
    style_header(ws, 1, len(cols))
    depr = {r["Sub_Category"]: r for r in read_csv(R_OUT / "depr_rates_demo.csv")}
    for cat in CATEGORIES:
        coefs = {r["Term"]: float(r["Estimate"])
                 for r in read_csv(R_OUT / f"coefficients_{cat.replace(' ', '_')}.csv")}
        usage = next((v for k, v in coefs.items() if "Km/" in k or "Hours/" in k), 0.0)
        ws.append([
            cat,
            round(coefs.get("(Intercept)", 0), 5),
            round(coefs.get("Age", 0), 5),
            round(usage, 5),
            round(coefs.get("LogNRC", 0), 5),
            round(coefs.get("Title_StatusRebuilt", 0), 5),
            float(depr[cat]["ResidualSE"]),
        ])
    for c in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions["A"].width = 20

    # --- Estimator --------------------------------------------------------
    ws = wb.create_sheet("Estimator")
    ws["A1"] = "Asset FMV Estimator"
    ws["A1"].font = TITLE
    ws["A3"], ws["B3"] = "Category", "Highway Tractor"
    ws["A4"], ws["B4"] = "Age (years)", 6
    ws["A5"], ws["B5"] = "Usage (Km for tractors / Hours for excavators)", 540000
    ws["A6"], ws["B6"] = "New Replacement Cost (NRC, $)", 195000
    ws["A7"], ws["B7"] = "Title status rebuilt? (0/1)", 0
    dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"')
    ws.add_data_validation(dv)
    dv.add(ws["B3"])
    # usage scale: Km/100000 for tractors, Hours/1000 for excavators, 0 for trailers
    ws["A9"], ws["B9"] = "Usage (scaled)", \
        '=IF(B3="Highway Tractor",B5/100000,IF(B3="Crawler Excavator",B5/1000,0))'
    ws["A10"], ws["B10"] = "log(FMV) prediction", (
        "=VLOOKUP(B3,Model_Params!A:G,2,FALSE)"
        "+VLOOKUP(B3,Model_Params!A:G,3,FALSE)*B4"
        "+VLOOKUP(B3,Model_Params!A:G,4,FALSE)*B9"
        "+VLOOKUP(B3,Model_Params!A:G,5,FALSE)*LN(B6)"
        "+VLOOKUP(B3,Model_Params!A:G,6,FALSE)*B7"
    )
    ws["A12"], ws["B12"] = "Estimated FMV", "=ROUND(EXP(B10),-2)"
    ws["A13"], ws["B13"] = "Low (80% band)", \
        "=ROUND(EXP(B10-1.282*VLOOKUP(B3,Model_Params!A:G,7,FALSE)),-2)"
    ws["A14"], ws["B14"] = "High (80% band)", \
        "=ROUND(EXP(B10+1.282*VLOOKUP(B3,Model_Params!A:G,7,FALSE)),-2)"
    for r in (12, 13, 14):
        ws.cell(row=r, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18

    # --- Depr_Rates --------------------------------------------------------
    ws = wb.create_sheet("Depr_Rates")
    rows = read_csv(R_OUT / "depr_rates_demo.csv")
    cols = list(rows[0].keys())
    ws.append(cols)
    style_header(ws, 1, len(cols))
    for r in rows:
        ws.append([r[c] for c in cols])
    for c in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16

    # --- Market_Comps ------------------------------------------------------
    ws = wb.create_sheet("Market_Comps")
    cols = ["Workfile", "Item", "FMV_Low", "FMV_High", "CompType", "CompText"]
    ws.append(cols)
    style_header(ws, 1, len(cols))
    for r in read_csv(EXTRACTED):
        for i in range(1, 7):
            if r.get(f"CompType{i}"):
                ws.append([r["Workfile"], r["Item"], r["FMV_Low"], r["FMV_High"],
                           r[f"CompType{i}"], r[f"CompText{i}"]])
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["F"].width = 80

    # --- Client_Summary (worksheet layout, fake client) --------------------
    ws = wb.create_sheet("Client_Summary")
    ws["A1"] = "Cedar Ridge Logistics Ltd. - Fleet Appraisal Summary (SYNTHETIC)"
    ws["A1"].font = TITLE
    ws["A2"] = "Verus Demo Appraisals"
    ws.append([])
    ws.append(["Asset Class", "# of units", "Book Value", "Demo FMV", "Demo OLV"])
    style_header(ws, 4, 5)
    for row in [("Highway Tractors", 14, 1284000, 1130000, 921000),
                ("Crawler Excavators", 6, 742000, 688000, 561000),
                ("Dry Van Trailers", 22, 814000, 759000, 619000)]:
        ws.append(row)
    ws.append(["TOTAL", "=SUM(B5:B7)", "=SUM(C5:C7)", "=SUM(D5:D7)", "=SUM(E5:E7)"])
    ws.cell(row=8, column=1).font = Font(bold=True)
    for c in range(1, 6):
        ws.column_dimensions[get_column_letter(c)].width = 18

    wb.save(OUT_XLSX)
    print(f"Dashboard written -> {OUT_XLSX}")


if __name__ == "__main__":
    main()
